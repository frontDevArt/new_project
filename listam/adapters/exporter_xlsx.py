"""Выгрузка в .xlsx: шапка заморожена, автофильтр, ссылки кликабельные,
цена и площадь — числа, сортировка по дате появления."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from listam.adapters.filenames import safe_filename
from listam.domain.models import Listing, Match, Request
from listam.ports.exporter import Exporter

SELLER_TYPES = {"owner": "собственник", "agency": "агентство"}
STATUSES = {"active": "на ленте", "gone": "снято"}
# Статусы матча по-русски: правило, добытое F-12 в M1 — в витрине не бывает
# английских слов из схемы базы.
MATCH_STATUSES = {"new": "новый", "sent": "отправлен",
                  "called": "звонили", "rejected": "отказ"}

# заголовок, как достать значение, ширина колонки, формат числа
COLUMNS: list[tuple[str, str, int, str | None]] = [
    ("ID", "id", 12, None),
    ("Ссылка", "url", 16, None),
    ("Заголовок", "title", 46, None),
    ("Район", "district", 18, None),
    ("Улица", "street", 22, None),
    ("Цена, $", "price_usd", 13, "#,##0"),
    ("Цена, ֏", "price_amd", 15, "#,##0"),
    ("Цена как на сайте", "price_raw", 18, None),
    ("Цена в валюте", "price_amount", 14, "#,##0.##"),
    ("Площадь, м²", "area", 12, "0.0"),
    ("$/м²", "price_per_sqm", 10, "#,##0"),
    ("Комнат", "rooms", 9, "0"),
    ("Этаж", "floor", 8, "0"),
    ("Этажей", "floors_total", 9, "0"),
    ("Продавец", "seller_type", 14, None),
    ("Новостройка", "new_build", 13, None),
    ("Проверено", "verified", 11, None),
    ("Сомнительно", "anomaly", 20, None),
    ("Статус", "status", 10, None),
    ("Появилось", "first_seen", 18, None),
    ("Видели", "last_seen", 18, None),
    ("Снято", "gone_at", 18, None),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINK_FONT = Font(color="0563C1", underline="single")


class XlsxExporter(Exporter):
    def __init__(self, directory: str | Path = "./out", timezone_name: str = "UTC"):
        self.directory = Path(directory)
        self.timezone_name = timezone_name

    def export(self, listings: Iterable[Listing], name: str | None = None,
               matches: list[tuple[Request, Match, Listing]] | None = None) -> Path:
        rows = sorted(
            listings,
            key=lambda item: (item.first_seen or datetime.min.replace(tzinfo=timezone.utc)),
            reverse=True,
        )
        self.directory.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
        # Имя приходит из командной строки: берём из него только имя файла.
        # `--name ../../отчёт.xlsx` обязан писать в свою папку, а не мимо неё.
        path = self.directory / (safe_filename(name) or f"listam-{stamp}.xlsx")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Объявления"

        for index, (title, _, width, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=1, column=index, value=title)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            sheet.column_dimensions[get_column_letter(index)].width = width

        for row_index, listing in enumerate(rows, start=2):
            for column_index, (title, attribute, _, number_format) in enumerate(COLUMNS, start=1):
                value = _present(listing, attribute)
                cell = sheet.cell(row=row_index, column=column_index, value=value)
                if number_format and isinstance(value, (int, float)):
                    cell.number_format = number_format
                if title == "Ссылка" and listing.url:
                    cell.hyperlink = listing.url
                    cell.font = LINK_FONT

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(sheet.max_row, 1)}"

        # Матчей нет — листа нет, и это не ошибка: выгрузка объявлений жила
        # без него всю M1 и обязана выглядеть ровно как раньше.
        if matches:
            _write_matches_sheet(workbook, list(matches))

        workbook.save(path)
        return path


def _present(listing: Listing, attribute: str):
    value = getattr(listing, attribute, None)
    if attribute == "url":
        return "открыть" if value else None
    if attribute == "seller_type":
        return SELLER_TYPES.get(value, value)
    if attribute == "status":
        return STATUSES.get(value, value)
    if attribute in ("new_build", "verified"):
        return None if value is None else ("да" if value else "нет")
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M")
    return value


# Лист «Матчи»: заголовок, откуда берётся значение, ширина, формат числа.
# Единица строки — кластер, а не объявление: «объявлений в кластере» и
# «разброс» лежат в самом матче снимком, сделанным подбором.
MATCH_COLUMNS: list[tuple[str, str, int, str | None]] = [
    ("Заявка", "request.external_id", 12, None),
    ("Клиент", "request.client_name", 18, None),
    ("Балл", "match.score", 8, "0"),
    ("Статус матча", "match.status", 14, None),
    ("Причина отказа", "match.reject_reason", 24, None),
    ("Цена, $", "listing.price_usd", 13, "#,##0"),
    ("$/м²", "listing.price_per_sqm", 10, "#,##0"),
    ("Район", "listing.district", 18, None),
    ("Улица", "listing.street", 22, None),
    ("Комнат", "listing.rooms", 9, "0"),
    ("Площадь, м²", "listing.area", 12, "0.0"),
    ("Этаж", "listing.floor", 8, "0"),
    ("Этажей", "listing.floors_total", 9, "0"),
    ("Продавец", "listing.seller_type", 14, None),
    ("Объявлений в кластере", "match.cluster_size", 22, "0"),
    ("Разброс, $", "match.cluster_spread_usd", 12, "#,##0"),
    ("Статус объявления", "listing.status", 18, None),
    ("Ссылка", "listing.url", 16, None),
]


def _match_value(request: Request, match: Match, listing: Listing, key: str):
    owner, attribute = key.split(".", 1)
    source = {"request": request, "match": match, "listing": listing}[owner]
    value = getattr(source, attribute, None)
    if owner == "match" and attribute == "status":
        return MATCH_STATUSES.get(value, value)
    if owner == "listing":
        return _present(listing, attribute)
    return value


def _write_matches_sheet(workbook: Workbook, rows: list[tuple]) -> None:
    """Лист витрины: те же правила, что на листе объявлений.

    Матч на снятое объявление остаётся в выгрузке с пометкой в колонке
    «Статус объявления» (решение 8): «мы звонили по этой квартире» не
    должно исчезать вместе с объявлением.
    """
    sheet = workbook.create_sheet("Матчи")
    for index, (title, _, width, _) in enumerate(MATCH_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        sheet.column_dimensions[get_column_letter(index)].width = width

    ordered = sorted(
        rows,
        key=lambda row: (row[0].external_id or "", -(row[1].score or 0.0),
                         row[2].id or ""),
    )
    for row_index, (request, match, listing) in enumerate(ordered, start=2):
        for column_index, (title, key, _, number_format) in enumerate(MATCH_COLUMNS,
                                                                      start=1):
            value = _match_value(request, match, listing, key)
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if number_format and isinstance(value, (int, float)):
                cell.number_format = number_format
            if title == "Ссылка" and listing.url:
                cell.hyperlink = listing.url
                cell.font = LINK_FONT

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(MATCH_COLUMNS))}{max(sheet.max_row, 1)}"
    )
