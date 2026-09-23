"""Контрактный тест порта Exporter: витрина для человека."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from listam.adapters.exporter_xlsx import XlsxExporter
from listam.domain.models import Listing, Match, Request
from listam.ports.exporter import Exporter

NOW = datetime(2026, 9, 21, 10, 0, tzinfo=timezone.utc)  # календарь: не сравнивается с часами


def listing(listing_id: str, first_seen: datetime, **over) -> Listing:
    fields = dict(
        id=listing_id,
        url=f"https://www.list.am/ru/item/{listing_id}",
        title="3 комн. квартира",
        district="Центр",
        price_raw="$132,000",
        currency="USD",
        price_usd=132000.0,
        price_amd=47916000.0,
        area=85.0,
        rooms=3,
        floor=4,
        floors_total=9,
        price_per_sqm=1552.94,
        seller_type="owner",
        verified=True,
        new_build=False,
        first_seen=first_seen,
        last_seen=first_seen,
    )
    fields.update(over)
    return Listing(**fields)


@pytest.fixture(params=["xlsx_local"])
def exporter(request, tmp_path) -> Exporter:
    return XlsxExporter(directory=tmp_path)


def test_is_an_exporter(exporter):
    assert isinstance(exporter, Exporter)


def test_writes_a_file_and_returns_its_path(exporter):
    path = exporter.export([listing("1", NOW)])
    assert path.exists()
    assert path.suffix == ".xlsx"


def test_header_is_frozen_and_filtered(exporter):
    sheet = load_workbook(exporter.export([listing("1", NOW)])).active
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


def test_price_and_area_are_numbers_not_text(exporter):
    sheet = load_workbook(exporter.export([listing("1", NOW)])).active
    header = [cell.value for cell in sheet[1]]
    row = {header[i]: cell.value for i, cell in enumerate(sheet[2])}
    assert isinstance(row["Цена, $"], (int, float))
    assert isinstance(row["Площадь, м²"], (int, float))
    assert isinstance(row["$/м²"], (int, float))


def test_url_is_a_clickable_link(exporter):
    sheet = load_workbook(exporter.export([listing("24254997", NOW)])).active
    header = [cell.value for cell in sheet[1]]
    cell = sheet[2][header.index("Ссылка")]
    assert cell.hyperlink.target == "https://www.list.am/ru/item/24254997"


def test_rows_are_sorted_by_first_seen_newest_first(exporter):
    older = listing("old", NOW - timedelta(days=2))
    newer = listing("new", NOW)
    sheet = load_workbook(exporter.export([older, newer])).active
    header = [cell.value for cell in sheet[1]]
    ids = [row[header.index("ID")].value for row in sheet.iter_rows(min_row=2)]
    assert ids == ["new", "old"]


def test_seller_type_is_written_in_russian(exporter):
    sheet = load_workbook(exporter.export([listing("1", NOW, seller_type="agency")])).active
    header = [cell.value for cell in sheet[1]]
    assert sheet[2][header.index("Продавец")].value == "агентство"


def test_empty_fields_do_not_break_the_export(exporter):
    sparse = listing("1", NOW, price_usd=None, area=None, rooms=None, district=None,
                     price_per_sqm=None, seller_type=None, verified=None)
    sheet = load_workbook(exporter.export([sparse])).active
    assert sheet.max_row == 2


def test_export_of_empty_database_still_writes_header(exporter):
    sheet = load_workbook(exporter.export([])).active
    assert sheet.max_row == 1
    assert sheet["A1"].value == "ID"


def test_export_has_a_column_for_suspicious_rows(tmp_path):
    from openpyxl import load_workbook

    from listam.adapters.exporter_xlsx import XlsxExporter
    from listam.domain.models import Listing, Match, Request

    rows = [
        Listing(id="1", url="u", district="Центр", area=1.0, rooms=3,
                anomaly="area,rooms_vs_area"),
        Listing(id="2", url="u", district="Центр", area=80.0, rooms=3),
    ]
    path = XlsxExporter(directory=tmp_path).export(rows, name="report.xlsx")

    sheet = load_workbook(path).active
    headers = [cell.value for cell in sheet[1]]
    column = headers.index("Сомнительно") + 1
    marks = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=column).value
             for r in (2, 3)}

    assert marks["1"] == "area,rooms_vs_area"
    assert marks["2"] is None


# --- M6: имя выгрузки не выводит из папки out/ ---

def test_name_cannot_lead_out_of_the_export_directory(tmp_path):
    out = tmp_path / "out"
    path = XlsxExporter(directory=out).export([listing("1", NOW)], name="../../беглец.xlsx")

    assert path.parent == out
    assert not (tmp_path.parent / "беглец.xlsx").exists()
    assert not (tmp_path / "беглец.xlsx").exists()


def test_name_without_a_filename_falls_back_to_the_default(tmp_path):
    path = XlsxExporter(directory=tmp_path).export([listing("1", NOW)], name="../")

    assert path.parent == tmp_path
    assert path.name.startswith("listam-")


def test_amount_in_original_currency_has_its_own_column(tmp_path):
    """ВЫСОКИЙ 9: у EUR пересчитанные колонки пусты, но число видно в «Цене в валюте»."""
    from openpyxl import load_workbook

    exporter = XlsxExporter(directory=tmp_path)
    path = exporter.export([
        Listing(id="1", url="https://www.list.am/ru/item/1", price_raw="140,000 €",
                currency="EUR", price_amount=140000.0, first_seen=NOW, last_seen=NOW),
    ])

    sheet = load_workbook(path).active
    headers = [cell.value for cell in sheet[1]]
    assert headers.index("Цена в валюте") == headers.index("Цена как на сайте") + 1
    assert len(headers) == 22
    assert sheet.auto_filter.ref.startswith("A1:V")
    assert sheet.cell(row=2, column=headers.index("Цена в валюте") + 1).value == 140000.0


def test_a_gone_listing_shows_its_status_and_the_day_it_left(exporter):
    """Снятое объявление остаётся в выгрузке: цена ушедшей квартиры — история рынка."""
    gone = Listing(id="1", url="https://www.list.am/ru/item/1", status="gone",
                   first_seen=NOW, last_seen=NOW,
                   gone_at=datetime(2026, 9, 22, 8, 0, tzinfo=timezone.utc))  # календарь: не сравнивается с часами

    sheet = load_workbook(exporter.export([gone])).active
    headers = [cell.value for cell in sheet[1]]

    assert headers[-1] == "Снято"
    assert sheet.cell(row=2, column=len(headers)).value == "2026-09-22 08:00"
    assert sheet.cell(row=2, column=headers.index("Статус") + 1).value == "снято"


def test_status_is_written_in_russian_like_its_neighbours(exporter):
    """«Продавец», «Новостройка» и «Проверено» переведены — «Статус» читается так же."""
    sheet = load_workbook(exporter.export([listing("1", NOW)])).active
    headers = [cell.value for cell in sheet[1]]

    assert sheet.cell(row=2, column=headers.index("Статус") + 1).value == "на ленте"


def test_the_status_column_has_no_latin_left_in_it(exporter):
    """Латиница в русской таблице — это утечка кода наружу, а не значение."""
    rows = [listing("1", NOW), Listing(id="2", url="https://www.list.am/ru/item/2",
                                       status="gone", first_seen=NOW, last_seen=NOW)]

    sheet = load_workbook(exporter.export(rows)).active
    headers = [cell.value for cell in sheet[1]]
    column = headers.index("Статус") + 1
    printed = [sheet.cell(row=r, column=column).value for r in (2, 3)]

    assert not any(any("a" <= ch.lower() <= "z" for ch in value) for value in printed)


# --- лист «Матчи» (фаза 6 M2) -----------------------------------------


def match_row(**over):
    """Строка витрины: заявка, матч и объявление вместе."""
    request = Request(id=1, external_id="R-1", client_name="Ани",
                      client_phone="+374 00 000000")
    fields = dict(id=1, request_id=1, listing_id="1", score=85.0, status="called",
                  reject_reason="окна во двор", cluster_id="c1", cluster_size=3,
                  cluster_spread_usd=11000.0)
    fields.update(over)
    return request, Match(**fields), listing("1", NOW)


def matches_sheet(exporter, rows):
    return load_workbook(exporter.export([listing("1", NOW)], matches=rows))["Матчи"]


def test_the_workbook_has_a_matches_sheet_when_matches_are_given(exporter):
    path = exporter.export([listing("1", NOW)], matches=[match_row()])

    book = load_workbook(path)
    assert "Матчи" in book.sheetnames
    sheet = book["Матчи"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref.startswith("A1")
    assert sheet.cell(row=2, column=1).value == "R-1"


def test_without_matches_the_workbook_is_exactly_what_it_was(exporter):
    book = load_workbook(exporter.export([listing("1", NOW)]))
    assert "Матчи" not in book.sheetnames
    assert book.sheetnames == ["Объявления"]


def test_the_matches_sheet_keeps_numbers_as_numbers(exporter):
    sheet = matches_sheet(exporter, [match_row()])
    values = {sheet.cell(row=1, column=i).value: sheet.cell(row=2, column=i).value
              for i in range(1, sheet.max_column + 1)}

    assert values["Балл"] == 85.0
    assert values["Цена, $"] == 132000.0
    assert values["Объявлений в кластере"] == 3
    assert values["Разброс, $"] == 11000.0


def test_the_match_status_is_russian_on_the_sheet(exporter):
    """Правило, добытое F-12 в M1: в витрине не бывает английских статусов."""
    sheet = matches_sheet(exporter, [match_row()])
    row = [sheet.cell(row=2, column=i).value for i in range(1, sheet.max_column + 1)]

    assert "звонили" in row
    assert "called" not in row


def test_a_match_on_a_listing_that_went_away_is_marked_and_not_hidden(exporter):
    """Решение 8: звонок переживает уход объявления с ленты."""
    gone = listing("1", NOW, status="gone", gone_at=NOW + timedelta(days=1))
    request, match, _ = match_row()

    sheet = matches_sheet(exporter, [(request, match, gone)])
    row = [sheet.cell(row=2, column=i).value for i in range(1, sheet.max_column + 1)]

    assert sheet.max_row == 2          # строка на месте, а не выброшена
    assert "снято" in row


def test_the_matches_sheet_has_no_none_written_out_as_a_word(exporter):
    """Слово `None` в клетке читается как поломка.

    Пустая клетка — не то же самое: в таблице она пуста и на листе
    объявлений, и человек видит именно пустоту, а не английское слово.
    """
    bare = listing("2", NOW, district=None, street=None, price_usd=None,
                   price_per_sqm=None, rooms=None, floor=None, floors_total=None,
                   seller_type=None)
    request, match, _ = match_row()

    sheet = matches_sheet(exporter, [(request, match, bare)])
    row = [sheet.cell(row=2, column=i).value for i in range(1, sheet.max_column + 1)]

    assert "None" not in [value for value in row if isinstance(value, str)]


def test_the_link_on_the_matches_sheet_is_clickable(exporter):
    sheet = matches_sheet(exporter, [match_row()])
    titles = [sheet.cell(row=1, column=i).value for i in range(1, sheet.max_column + 1)]
    cell = sheet.cell(row=2, column=titles.index("Ссылка") + 1)

    assert cell.hyperlink is not None
